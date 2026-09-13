"""Pruebas de los reportes descargables y del rendimiento de la consulta."""

from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import event, update

from app.consultas import ZONA_LOCAL
from app.consultas import movimientos as consulta_movimientos
from app.models import AuditoriaNegocio


def _mover(cliente, encabezados, qr, tipo):
    return cliente.post("/movimientos", headers=encabezados, json={"qr": qr, "tipo": tipo})


def test_reporte_excel_contiene_los_movimientos(cliente, dispositivo, encabezados_admin):
    _mover(cliente, encabezados_admin, dispositivo.qr, "Ingreso")

    respuesta = cliente.get("/reportes/movimientos.xlsx", headers=encabezados_admin)
    assert respuesta.status_code == 200
    assert "attachment" in respuesta.headers["content-disposition"]

    hoja = load_workbook(BytesIO(respuesta.content)).active
    assert hoja["A1"].value == "Fecha"
    assert hoja["C2"].value == dispositivo.serial
    assert hoja["B2"].value == "Ingreso"


def test_reporte_pdf_es_un_pdf_valido(cliente, dispositivo, encabezados_admin):
    _mover(cliente, encabezados_admin, dispositivo.qr, "Ingreso")

    respuesta = cliente.get("/reportes/movimientos.pdf", headers=encabezados_admin)
    assert respuesta.status_code == 200
    assert respuesta.headers["content-type"] == "application/pdf"
    assert respuesta.content.startswith(b"%PDF-")


def test_reporte_vacio_no_falla(cliente, encabezados_admin):
    """Sin movimientos el reporte debe generarse igual, no reventar."""
    assert cliente.get("/reportes/movimientos.pdf", headers=encabezados_admin).status_code == 200
    assert cliente.get("/reportes/movimientos.xlsx", headers=encabezados_admin).status_code == 200


def test_el_reporte_respeta_el_filtro_por_tipo(
    cliente, dispositivo, otro_dispositivo, encabezados_admin
):
    _mover(cliente, encabezados_admin, dispositivo.qr, "Ingreso")
    _mover(cliente, encabezados_admin, dispositivo.qr, "Salida")
    _mover(cliente, encabezados_admin, otro_dispositivo.qr, "Ingreso")

    respuesta = cliente.get("/reportes/movimientos.xlsx?tipo=Salida", headers=encabezados_admin)
    hoja = load_workbook(BytesIO(respuesta.content)).active

    # Encabezado + una sola fila de datos.
    assert hoja.max_row == 2
    assert hoja["B2"].value == "Salida"


def test_el_filtro_de_fecha_es_inclusivo(cliente, dispositivo, encabezados_admin):
    """Filtrar 'hasta hoy' debe incluir los movimientos de hoy."""
    from datetime import date

    _mover(cliente, encabezados_admin, dispositivo.qr, "Ingreso")
    hoy = date.today().isoformat()

    movimientos = cliente.get(
        f"/movimientos?desde={hoy}&hasta={hoy}", headers=encabezados_admin
    ).json()
    assert len(movimientos) == 1


def test_el_filtro_incluye_el_turno_de_la_noche(cliente, db, dispositivo, encabezados_admin):
    """Un movimiento de las 11 de la noche pertenece a ese día, no al siguiente.

    Regresión: los límites del filtro se construían sin zona horaria y Postgres
    los interpretaba como UTC. En Colombia (UTC-5) eso mandaba todo lo
    registrado después de las 19:00 al día siguiente, así que el turno de la
    tarde desaparecía del filtro de su propio día.

    La fecha se fija a mano en vez de usar `date.today()`: si dependiera del
    reloj, la prueba solo fallaría al correrla de noche — que es justamente cómo
    este defecto sobrevivió a un CI en verde, porque los runners van en UTC.
    """
    _mover(cliente, encabezados_admin, dispositivo.qr, "Ingreso")
    db.execute(
        update(AuditoriaNegocio).values(
            fecha_creado=datetime(2026, 8, 11, 23, 30, tzinfo=ZONA_LOCAL)
        )
    )
    db.commit()

    movimientos = cliente.get(
        "/movimientos?desde=2026-08-11&hasta=2026-08-11", headers=encabezados_admin
    ).json()
    assert len(movimientos) == 1


def test_la_consulta_no_dispara_una_avalancha_de_queries(db, dispositivo, otro_dispositivo, admin):
    """Regresión de N+1: leer 4 movimientos debe costar una consulta, no 17."""
    from app.porteria import TipoMovimiento, registrar

    for equipo in (dispositivo, otro_dispositivo):
        registrar(db, equipo, TipoMovimiento.INGRESO, admin.id)
        registrar(db, equipo, TipoMovimiento.SALIDA, admin.id)

    ejecutadas: list[str] = []

    def contar(conn, cursor, sentencia, *args):
        ejecutadas.append(sentencia)

    event.listen(db.get_bind(), "before_cursor_execute", contar)
    try:
        resultados = db.scalars(consulta_movimientos()).unique().all()
        # Tocar las relaciones: si no vinieran precargadas, aquí saltarían las
        # consultas adicionales.
        for m in resultados:
            _ = (
                m.dispositivo.serial,
                m.dispositivo.responsable,
                m.tipo_registro.nombre,
                m.guarda.nombre_completo,
            )
    finally:
        event.remove(db.get_bind(), "before_cursor_execute", contar)

    assert len(resultados) == 4
    assert len(ejecutadas) == 1, f"se esperaba 1 consulta, se ejecutaron {len(ejecutadas)}"
